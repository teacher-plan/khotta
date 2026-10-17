# كتابة ورقة الاختبار ودليل التصحيح بصيغة Word (.docx) من اليمين لليسار، مع كسور حقيقية (معادلات Word).
# تحتاج: pip install python-docx
import os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
import notation
from exam_lib import (plain, ar, marks_word, LETTERS, TYPE_NAMES, tool_cfg, key_columns, field_value,
                      grid_model, summary_model, KEY_WIDTHS, item_labels)

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

# ترتيب العناصر كما في مواصفة OOXML (Word قد يرفض الملف إن خالفه)
SEQ = {
    'w:pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs '
             'suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd '
             'snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment '
             'textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
    'w:rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid '
             'vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em '
             'lang eastAsianLayout specVanish oMath',
    'w:tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing '
               'tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription',
    'w:tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark',
    'w:sectPr': 'headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType '
                'pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange',
}


def fix_order(root):
    for tag, seq in SEQ.items():
        order = {qn('w:' + n): i for i, n in enumerate(seq.split())}
        for el in root.iter(qn(tag)):
            kids = list(el)
            seen = set(); uniq = []
            for k in kids:   # عنصر مفرد مكرّر (مثل szCs) يبقى آخره
                if k.tag in seen and k.tag != qn('w:headerReference') and k.tag != qn('w:footerReference'):
                    uniq = [u for u in uniq if u.tag != k.tag]
                seen.add(k.tag); uniq.append(k)
            for k in kids:
                el.remove(k)
            for k in sorted(uniq, key=lambda k: order.get(k.tag, 999)):
                el.append(k)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


class DocxExam:
    def __init__(self, font='Arial', digits='ar', snip=None, quran_font='Traditional Arabic'):
        self.font, self.dg, self.snip, self.quran_font = font, digits, snip, quran_font

    @staticmethod
    def save(d, path):
        fix_order(d.element)
        fix_order(d.styles.element)
        d.save(path)

    # ── أساسيات ──
    def new_doc(self, landscape=False):
        d = Document()
        s = d.sections[0]
        s.page_width, s.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
        s.left_margin = s.right_margin = Mm(13); s.top_margin = Mm(12); s.bottom_margin = Mm(12)
        s._sectPr.append(_el('w:bidi'))
        st = d.styles['Normal']
        st.font.name = self.font; st.font.size = Pt(12.5)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = _el('w:rFonts'); rpr.append(rf)
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(k), self.font)
        rpr.append(_el('w:szCs', **{'w:val': 25}))
        rpr.append(_el('w:lang', **{'w:bidi': 'ar-OM'}))
        st.paragraph_format.space_after = Pt(2)
        return d

    def para(self, cont, align='right', space_before=0, space_after=2, keep=False):
        p = cont.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(space_before), Pt(space_after)
        if keep:
            pf.keep_with_next = True
        p._p.get_or_add_pPr().insert(0, _el('w:bidi'))
        # في الفقرة العربية (bidi) يقرأ Word وLibreOffice القيمة left على أنها «البداية» أي اليمين، وright على أنها اليسار
        p.alignment = {'right': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'left': WD_ALIGN_PARAGRAPH.RIGHT}[align]
        return p

    def run(self, p, text, bold=False, size=None, color=None, sup=False, sub=False, rtl=True, font=None):
        r = p.add_run(text)
        rpr = r._r.get_or_add_rPr()
        if rtl:
            rpr.append(_el('w:rtl'))
        if font:
            rpr.append(_el('w:rFonts', **{'w:ascii': font, 'w:hAnsi': font, 'w:cs': font}))
        if bold:
            r.bold = True; rpr.append(_el('w:bCs'))
        if size:
            r.font.size = Pt(size); rpr.append(_el('w:szCs', **{'w:val': int(size * 2)}))
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if sup:
            r.font.superscript = True
        if sub:
            r.font.subscript = True
        return r

    def _mrun(self, text, bold=False):
        r = OxmlElement('m:r')
        rp = OxmlElement('m:rPr'); rp.append(_el('m:sty', **{'m:val': 'b' if bold else 'p'})); r.append(rp)
        w = OxmlElement('w:rPr'); w.append(_el('w:rFonts', **{'w:ascii': self.font, 'w:hAnsi': self.font, 'w:cs': self.font}))
        r.append(w)
        t = OxmlElement('m:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
        return r

    def omath(self, p, elements):
        om = OxmlElement('m:oMath')
        for e in elements:
            om.append(e)
        p._p.append(om)

    def rich(self, p, text, bold=False, size=None):
        self.nodes(p, notation.parse(text, self.dg), bold, size)
        return p

    def nodes(self, p, ns, bold=False, size=None, sup=False, sub=False, rtl=True, under=False):
        mk = notation.OMML(lambda tag, **kw: _el(tag, **kw), lambda t: self._mrun(t, bold))
        for n in ns:
            k = n[0]
            if k == 't':
                r_ = self.run(p, n[1], bold, size, sup=sup, sub=sub, rtl=rtl)
                if under:
                    r_.underline = True
            elif k == 'u':
                self.nodes(p, n[1], bold, size, sup, sub, rtl, under=True)
            elif k == 'b':
                self.nodes(p, n[1], True, size, sup, sub, rtl)
            elif k == 'grp':
                self.nodes(p, n[1], bold, size, sup, sub, rtl)
            elif k in ('sup', 'sub'):
                # في النص العربي: علامة RLM تفصل أرقام الأسّ عن أرقام الأساس فيظهر الأسّ يسار الأساس كما في الكتاب؛
                # وفي النص اللاتيني (العلوم) يبقى الأسّ يمين أساسه
                if rtl:
                    self.run(p, notation.RLM, size=size)
                self.nodes(p, n[1], bold, size, sup=(k == 'sup'), sub=(k == 'sub'), rtl=rtl)
            elif k == 'frac':
                if n[1]:   # العدد الصحيح نصّاً قبل الكسر (يمينه)
                    self.nodes(p, n[1], bold, size)
                self.omath(p, mk.atom(('node', ('frac', None, n[2], n[3]))))
            elif k == 'ol':   # خطّ فوق كل رقم (U+0305) في النص نفسه، فيبقى العدد متّصلاً
                self.run(p, ''.join(ch + '\u0305' for ch in notation.plain_nodes(n[1])), bold, size, sup=sup, sub=sub)
            elif k == 'root':
                img = self.snip(n) if self.snip else None
                if img:
                    path, w_pt, h_pt = img
                    scale = (size or 12.5) / 12.5
                    r = p.add_run(); r.add_picture(path, width=Pt(w_pt * scale), height=Pt(h_pt * scale))
                    # إنزال الصورة ليتوسّط الجذرُ السطرَ (w:position بأنصاف النقاط)
                    r._r.get_or_add_rPr().append(_el('w:position', **{'w:val': -int(h_pt * scale * 0.55)}))
                else:
                    self.omath(p, mk.atom(('node', n)))
            elif k in ('ce', 'ltr'):
                self.run(p, notation.LRE, size=size, rtl=False)
                if k == 'ce':
                    for t, kk in notation.ce_tokens(n[1]):
                        self.run(p, t, bold, size, sup=(kk == 'sup'), sub=(kk == 'sub'), rtl=False)
                else:
                    self.nodes(p, n[1], bold, size, sup, sub, rtl=False)
                self.run(p, notation.PDF_, size=size, rtl=False)
            elif k == 'quran':
                self.run(p, '﴿' + n[1].strip() + '﴾', bold, (size or 12.5) + 1.5, font=self.quran_font)
            elif k == 'blank':
                self.run(p, ' .................... ', False, size)
            elif k == 'br':
                p.add_run().add_break()

    def table(self, cont, rows, cols, widths=None, border=True, center=True):
        t = cont.add_table(rows=rows, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tpr = t._tbl.tblPr
        tpr.append(_el('w:bidiVisual'))
        if border:
            b = OxmlElement('w:tblBorders')
            for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                b.append(_el(f'w:{side}', **{'w:val': 'single', 'w:sz': 8, 'w:color': '333333'}))
            tpr.append(b)
        if widths:   # عرض ثابت للأعمدة (Word و LibreOffice يقرآن tblGrid)
            t.autofit = False
            tpr.append(_el('w:tblLayout', **{'w:type': 'fixed'}))
            for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
                gc.set(qn('w:w'), str(int(w * 56.7)))
            for row in t.rows:
                for c, w in zip(row.cells, widths):
                    c.width = Mm(w)
        return t

    def cell(self, c, text, bold=False, size=None, align='center', shade=None):
        c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
        for k, line in enumerate(str(text).split('\n')):
            p = self.para(c, align, 0, 0)
            self.rich(p, line, bold, size)
        if shade:
            tcpr = c._tc.get_or_add_tcPr(); tcpr.append(_el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': shade}))

    def lines(self, cont, n, width_mm=182):
        # سطر إجابة: نقاط نصّية بعرض الحاوية (علامة الجدولة المنقّطة لا تظهر داخل خلايا الجداول، وحدود الفقرات تُدمج)
        dots = '.' * int(width_mm / 1.13 * 0.9)   # النقطة بحجم ١١ ≈ ١٫١ مم في Arial
        for _ in range(n):
            p = self.para(cont, 'right', 9, 0)
            self.run(p, dots, size=11, color='777777')

    def header(self, d, hl, landscape=False):
        right, mid, left = hl
        w = (86, 99, 86) if landscape else (60, 64, 60)
        t = self.table(d, 1, 3, w, border=True)
        cells = t.rows[0].cells
        for c, xs, al in ((cells[0], right, 'right'), (cells[1], mid, 'center'), (cells[2], left, 'left')):
            c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
            for i, x in enumerate(xs):
                p = self.para(c, al, 0, 0)
                self.rich(p, x, bold=(c is cells[1] and i == 0) or x.startswith('النموذج'), size=14 if (c is cells[1] and i == 0) else 10.5)

    # ── ورقة الأسئلة ──
    def paper(self, exam, cfg, form, items, base, hl, calc):
        d = self.new_doc()
        self.header(d, hl)
        self.para(d, 'right', 0, 0).paragraph_format.line_spacing = Pt(6)
        t = self.table(d, 1, 2, (150, 34))
        c0, c1 = t.rows[0].cells
        self.cell(c0, 'اسم الطالب: ..............................................   الصف: ........', align='right', size=11.5)
        self.cell(c1, f'الدرجة\n      / {ar(exam["total"])}', bold=True, size=12)
        inst = '؛ '.join(x for x in [exam.get('instructions', 'أجب عن جميع الأسئلة في الورقة نفسها'), calc] if x) + '.'
        p = self.para(d, 'right', 4, 4); self.run(p, inst, size=11)
        labels = iter(item_labels(items)); grp = None
        for it in items:
            if it.get('group') and it['group'] != grp:
                grp = it['group']; p = self.para(d, 'right', 8, 2, keep=True); self.run(p, f'السؤال {grp}', True, 13.5)
            if it.get('type') == 'passage':
                tb = self.table(d, 1, 1, (184,))
                c = tb.rows[0].cells[0]; c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
                if it.get('title'):
                    self.rich(self.para(c, 'center', 2, 2), it['title'], bold=True)
                for line in (str(it['body']).split('\n') if it.get('body') else []):
                    self.rich(self.para(c, 'right', 1, 1), line)
                if it.get('verses'):
                    vt = self.table(c, len(it['verses']), 3, (78, 14, 78), border=False)
                    for r_, (a_, b_) in zip(vt.rows, it['verses']):
                        self.cell(r_.cells[0], a_, align='right'); self.cell(r_.cells[1], ' '); self.cell(r_.cells[2], b_, align='right')
                if it.get('source'):
                    self.rich(self.para(c, 'left', 2, 0), it['source'], size=10)
                continue
            self.item(d, next(labels)[0], it, base)
        p = self.para(d, 'center', 4, 0); self.run(p, 'انتهت الأسئلة، مع تمنياتنا لكم بالتوفيق', True, 11.5)
        if cfg.get('teacher_name'):
            self.run(p, f'          معلّم المادّة: {cfg["teacher_name"]}', size=10)
        return d

    # ── نمط «الجدول» (الافتراضي): كعيّنة مدارس سلطنة عُمان ──
    @staticmethod
    def borders(c, sz=12, val='single', sides=('top', 'left', 'bottom', 'right')):
        tcpr = c._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            b.append(_el(f'w:{side}', **({'w:val': val, 'w:sz': sz, 'w:color': '000000'} if side in sides else {'w:val': 'nil'})))
        tcpr.append(b)

    @staticmethod
    def vcenter(c):
        c._tc.get_or_add_tcPr().append(_el('w:vAlign', **{'w:val': 'center'}))

    def paper_table(self, exam, info, items, base):
        d = self.new_doc()
        st = d.styles['Normal']; st.font.size = Pt(14)   # كالعيّنة: الأسئلة ١٤، الجهات ١٢ عريض، العنوان ١٦ عريض
        st.element.get_or_add_rPr().find(qn('w:szCs')).set(qn('w:val'), '28')
        sp = d.sections[0]._sectPr   # إطار مزدوج حول كل صفحة
        pb = _el('w:pgBorders', **{'w:offsetFrom': 'page'})
        for side in ('top', 'left', 'bottom', 'right'):
            pb.append(_el(f'w:{side}', **{'w:val': 'double', 'w:sz': 12, 'w:space': 20, 'w:color': '000000'}))
        sp.append(pb)
        # الترويسة: الجهات يميناً، الشعار وسطاً، خانة الدرجة يساراً
        t = self.table(d, 1, 3, (104, 36, 44), border=False)
        cr, cm, cl = t.rows[0].cells
        cr.paragraphs[0]._p.getparent().remove(cr.paragraphs[0]._p)
        for x in info['right']:
            self.run(self.para(cr, 'center', 0, 0), x, True, 12)
        pm = cm.paragraphs[0]; pm.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if info.get('logo'):
            pm.add_run().add_picture(info['logo'], height=Mm(21))
        cl.paragraphs[0]._p.getparent().remove(cl.paragraphs[0]._p)
        sc = self.table(cl, 2, 1, (24,))
        sc.alignment = WD_TABLE_ALIGNMENT.RIGHT   # في الجدول المعكوس (bidiVisual) تعني اليسار
        for r_, v in zip(sc.rows, ('', info['total'])):
            r_.height = Mm(10)
            self.borders(r_.cells[0], 16); self.vcenter(r_.cells[0])
            self.cell(r_.cells[0], v or ' ', True, 13)
        self.para(cl, 'left', 0, 0)
        # العنوان مسطّراً
        for i, x in enumerate(info['title']):
            p = self.para(d, 'center', 6 if i == 0 else 0, 0)
            r_ = self.run(p, x, True, 16); r_.underline = True
        if info.get('meta'):
            self.run(self.para(d, 'center', 2, 0), '          '.join(info['meta']), size=12)
        # الاسم والصف
        self.para(d, 'right', 0, 0).paragraph_format.line_spacing = Pt(8)
        t = self.table(d, 1, 3, (130, 8, 46), border=False)
        a_, g_, b_ = t.rows[0].cells
        for c, x in ((a_, 'الاسم: ......................................................................'), (b_, info['class'])):
            self.borders(c, 12); self.vcenter(c); self.cell(c, x, size=12, align='right')
        t.rows[0].height = Mm(12)
        self.run(self.para(d, 'right', 6, 3), info['inst'], size=13)
        # جدول المفردات
        Wq = 144
        n_rows = 1 + len(items) + len({i.get('group') for i in items if i.get('group')})
        qt = self.table(d, n_rows, 3, (21, Wq, 19), border=False)
        b = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            b.append(_el(f'w:{side}', **{'w:val': 'double', 'w:sz': 12, 'w:color': '000000'}))
        for side in ('insideH', 'insideV'):
            b.append(_el(f'w:{side}', **{'w:val': 'single', 'w:sz': 12, 'w:color': '000000'}))
        qt._tbl.tblPr.append(b)
        hd = qt.rows[0]
        hd._tr.get_or_add_trPr().append(_el('w:tblHeader'))
        for c, x, al in zip(hd.cells, ('المفردة', 'السؤال', 'الدرجة'), ('center', 'right', 'center')):
            self.cell(c, x, True, 11.5, align=al); self.vcenter(c)
        ri = 1; grp = None
        labels = iter(item_labels(items))
        for it in items:
            if it.get('group') and it['group'] != grp:
                grp = it['group']; row = qt.rows[ri]; ri += 1
                m = row.cells[0].merge(row.cells[2]); self.cell(m, f'السؤال {grp}', True, 13, align='right', shade='ECECEC')
                m.paragraphs[0].paragraph_format.keep_with_next = True   # لا يبقى العنوان وحده آخر الصفحة
                row._tr.get_or_add_trPr().append(_el('w:cantSplit'))
            row = qt.rows[ri]; ri += 1
            row._tr.get_or_add_trPr().append(_el('w:cantSplit'))
            cn, cb, cmk = row.cells
            for c in (cn, cmk):
                self.vcenter(c)
            if it.get('type') == 'passage':
                self.cell(cn, ' '); self.cell(cmk, ' ')
                cb.paragraphs[0]._p.getparent().remove(cb.paragraphs[0]._p)
                if it.get('title'):
                    self.rich(self.para(cb, 'center', 2, 2), it['title'], bold=True)
                for line in (str(it['body']).split('\n') if it.get('body') else []):
                    self.rich(self.para(cb, 'right', 1, 1), line)
                if it.get('verses'):
                    vt = self.table(cb, len(it['verses']), 3, (66, 10, 66), border=False)
                    for r_, (x_, y_) in zip(vt.rows, it['verses']):
                        self.cell(r_.cells[0], x_, align='right'); self.cell(r_.cells[1], ' '); self.cell(r_.cells[2], y_, align='right')
                if it.get('source'):
                    self.rich(self.para(cb, 'left', 2, 0), it['source'], size=10)
                continue
            self.cell(cn, next(labels)[0], size=13)
            self.cell(cmk, f'[{ar(it["marks"])}]', size=13)
            first = cb.paragraphs[0]._p
            self.item(cb, None, it, base, W=Wq - 6, numbered=False, letters=info.get('letters', False))
            first.getparent().remove(first)
            self.para(cb, 'right', 0, 2)
        if info.get('closing'):
            self.para(d, 'right', 0, 0).paragraph_format.line_spacing = Pt(6)
            ct = self.table(d, 1, 1, (184,))
            ct.rows[0].height = Mm(13); self.vcenter(ct.rows[0].cells[0])
            self.cell(ct.rows[0].cells[0], info['closing'], size=13.5)
        return d

    def item(self, d, k, it, base, W=184, numbered=True, letters=True):
        t = it['type']
        p = self.para(d, 'right', 4, 2, keep=True)
        if numbered:
            self.run(p, f'{k}) ' if it.get('group') is None else f'({k}) ', True, 13.5)
        self.rich(p, it['stem'], bold=numbered)
        if numbered:
            self.run(p, f'   ({marks_word(it["marks"])})', size=10, color='444444')
        if it.get('fig'):
            path = it['fig'] if os.path.isabs(it['fig']) else os.path.join(base, it['fig'])
            pp = self.para(d, 'center', 2, 2)
            pp.add_run().add_picture(path, width=Mm(min(W - 6, 1.7 * it.get('fig_width', 45))))
        if it.get('graph') and getattr(self, 'graph_img', None):   # شبكة الإحداثيات صورةً بمقاسها الحقيقي
            gi = self.graph_img(self.graph_svg(it))
            if gi:
                pp = self.para(d, 'center', 2, 2)
                sc = min(1.0, Mm(W - 6).pt / gi[1])
                pp.add_run().add_picture(gi[0], width=Pt(gi[1] * sc), height=Pt(gi[2] * sc))
        if t == 'mcq':
            ch = it['choices']
            cols = it.get('choice_cols', 4 if max(len(plain(c)) for c in ch) < 14 else 2)
            rows = (len(ch) + cols - 1) // cols
            tb = self.table(d, rows, cols, [W / cols] * cols, border=numbered)
            for i, c in enumerate(ch):
                self.cell(tb.rows[i // cols].cells[i % cols], f'({LETTERS[i]})  {c}' if letters else c,
                          align='right' if numbered else 'center')
        elif t == 'tf':
            p = self.para(d, 'right', 2, 2); self.run(p, '(  صواب  /  خطأ  )', True)
            if it.get('explain', True):
                p = self.para(d, 'right', 2, 0); self.run(p, 'فسّر إجابتك:')
                self.lines(d, it.get('lines', 2), W - 4)
        elif t == 'match':
            L, R = it['left'], it['right']
            n = max(len(L), len(R))
            tb = self.table(d, n + 1, 5, [w * W / 164 for w in (12, 62, 16, 12, 62)])
            hdr = tb.rows[0].cells
            self.cell(hdr[0], ''); self.cell(hdr[1], it.get('left_title', 'العمود الأول'), True, shade='E6E6E6')
            self.cell(hdr[2], 'الحرف', True, shade='E6E6E6'); self.cell(hdr[3], '')
            self.cell(hdr[4], it.get('right_title', 'العمود الثاني'), True, shade='E6E6E6')
            for i in range(n):
                r = tb.rows[i + 1].cells
                if i < len(L):
                    self.cell(r[0], ar(i + 1)); self.cell(r[1], L[i]); self.cell(r[2], '')
                if i < len(R):
                    self.cell(r[3], LETTERS[i]); self.cell(r[4], R[i])
        elif t == 'order':
            p = self.para(d, 'right', 2, 2)
            self.rich(p, '     |     '.join(it['items']))
            p = self.para(d, 'right', 2, 0); self.run(p, 'الترتيب:')
            self.lines(d, it.get('lines', 1), W - 4)
        elif t == 'table':
            rows = it['grid']['rows']
            tb = self.table(d, len(rows), len(rows[0]))
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    hd = r == 0 and it['grid'].get('header', True)
                    self.cell(tb.rows[r].cells[c], v or ' ', hd, shade='E6E6E6' if hd else None)
            self.lines(d, it.get('lines', 0), W - 4)
        else:
            default = 0 if it.get('parts') else {'long': 6, 'explain': 3, 'creative': 4, 'short': 1, 'fill': 0}.get(t, 2)
            self.lines(d, it.get('lines', default), W - 4)
        for prt in it.get('parts', []):
            p = self.para(d, 'right', 4, 0); self.run(p, f'({prt["label"]}) ', True); self.rich(p, prt['stem'])
            self.lines(d, prt.get('lines', 2), W - 4)

    # ── دليل التصحيح ──
    def key(self, exam, cfg, spec, hl):
        tool = tool_cfg(spec, exam['kind'], exam['grade'], exam.get('number'))
        d = self.new_doc(landscape=True)
        self.header(d, hl, landscape=True)
        lessons = '، '.join(f'{l["code"]} {l.get("title", "")}' if isinstance(l, dict) else l for l in exam['lessons'])
        p = self.para(d, 'right', 4, 4)
        self.run(p, f'الدروس: {lessons} · المرجع: {spec["document"]["title"]}، {spec["document"]["edition"]}', size=10)
        cols = key_columns(spec, tool)
        tw = sum(KEY_WIDTHS.get(f, 9) for _, f in cols)
        for n, (form, items) in enumerate(exam['forms'].items()):
            items = [i for i in items if i.get('type') != 'passage']
            if n:
                d.add_page_break()
            p = self.para(d, 'right', 6, 2); self.run(p, f'شبكة مفردات النموذج ({form})', True, 13)
            heads, rows = grid_model(spec, items)
            tb = self.table(d, len(rows) + 1, len(heads), [270 / len(heads)] * len(heads))
            for c, v in zip(tb.rows[0].cells, heads):
                self.cell(c, v, True, 10.5, shade='E6E6E6')
            for k, r in enumerate(rows, 1):
                for c, v in zip(tb.rows[k].cells, r):
                    self.cell(c, v, size=10.5)
            sm = summary_model(spec, tool, items)
            self.para(d, 'right', 4, 0)
            tb = self.table(d, 2, len(sm), [270 / len(sm)] * len(sm))
            for c, (lab, v, w) in zip(tb.rows[0].cells, sm):
                self.cell(c, lab, True, 10, shade='E6E6E6')
            for c, (lab, v, w) in zip(tb.rows[1].cells, sm):
                self.cell(c, ar(v) + (f' / {ar(w)}' if w is not None else ''), size=10)
            p = self.para(d, 'right', 8, 2); self.run(p, f'دليل التصحيح — النموذج ({form})', True, 13)
            fl = [f for _, f in cols]
            tb = self.table(d, len(items) + 2, len(cols), [270 * KEY_WIDTHS.get(f, 9) / tw for f in fl])
            trpr = tb.rows[0]._tr.get_or_add_trPr(); trpr.append(_el('w:tblHeader'))   # تكرار العناوين في كل صفحة
            for c, (hd, _) in zip(tb.rows[0].cells, cols):
                self.cell(c, hd, True, 10.5, shade='E6E6E6')
            for k, (lab, it) in enumerate(zip(item_labels(items), items), 1):
                for c, f in zip(tb.rows[k].cells, fl):
                    if f == 'n':
                        v = lab[1]
                    elif f == 'marks':
                        v = ar(it.get('marks', 0))
                    elif f == 'answer':
                        v = self.answer_text(it)
                    elif f in ('objective', 'page', 'guidance'):
                        v = it.get(f, '')
                    else:
                        v = str(field_value(spec, it, f, exam))
                    self.cell(c, v or ' ', size=10.5, align='right' if f in ('objective', 'answer', 'guidance') else 'center')
                    if f == 'answer' and it.get('graph_answer') and getattr(self, 'graph_img', None):
                        gi = self.graph_img(self.graph_svg(it, True))
                        if gi:
                            pp = self.para(c, 'center', 2, 2)
                            sc = min(1.0, c.width.pt / gi[1] * 0.95) if c.width else 1.0
                            pp.add_run().add_picture(gi[0], width=Pt(gi[1] * sc), height=Pt(gi[2] * sc))
            last = tb.rows[-1].cells
            if 'marks' in fl:
                mi = fl.index('marks')
                if mi > 0:
                    m = last[0].merge(last[mi - 1]); self.cell(m, 'المجموع', True, 10.5, shade='E6E6E6')
                self.cell(last[mi], ar(sum(i['marks'] for i in items)), True, 10.5)
                if mi + 1 < len(fl):
                    m2 = last[mi + 1].merge(last[-1]) if mi + 1 < len(fl) - 1 else last[mi + 1]
                    self.cell(m2, ' ')
        p = self.para(d, 'right', 14, 0)
        self.run(p, f'إعداد: {cfg.get("teacher_name", "........................")}          '
                    f'مراجعة {cfg.get("reviewer_title", "المعلّم الأول")}: ........................   التوقيع: ............', size=11)
        return d

    def answer_text(self, it):
        if it['type'] == 'mcq':
            a = it['answer']; return (f'({LETTERS[a]}) ' if getattr(self, 'letters', True) else '') + it['choices'][a]
        if it['type'] == 'match' and it.get('pairs'):
            return '، '.join(f'{ar(i + 1)} ← {LETTERS[j]}' for i, j in enumerate(it['pairs']))
        parts = [str(it['answer'])] if it.get('answer') not in (None, '') else []
        parts += it.get('steps', [])
        return '\n'.join(parts)
